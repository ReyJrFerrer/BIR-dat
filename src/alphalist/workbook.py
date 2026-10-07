"""Bounded XLSX ingestion, matched by header text rather than column position."""

import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .domain import FilingContext, SourceRow, SourceWorkbook

HEADERS: dict[str, str] = json.loads(Path(__file__).with_name("headers.json").read_text())
MAX_UPLOAD = 10 * 1024 * 1024
MAX_EXPANDED = 40 * 1024 * 1024


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime | date):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def read_workbook(data: bytes, filename: str) -> SourceWorkbook:
    if len(data) > MAX_UPLOAD or not filename.lower().endswith(".xlsx"):
        raise ValueError("Choose an XLSX file smaller than 10 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(x.file_size for x in archive.infolist()) > MAX_EXPANDED
            ):
                raise ValueError("Workbook expands beyond the supported size limit.")
            if any(x.flag_bits & 1 for x in archive.infolist()):
                raise ValueError("Encrypted workbooks are not supported.")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("This file is not a readable XLSX workbook.") from exc
    candidates = []
    try:
        for sheet in book:
            if (
                sheet.max_row
                and sheet.max_row > 10000
                or sheet.max_column
                and sheet.max_column > 150
            ):
                raise ValueError("Supported workbooks have at most 10,000 rows and 150 columns.")
            rows = list(sheet.iter_rows())
            for index, row in enumerate(rows[:100]):
                labels = [text(c.value) for c in row]
                if HEADERS["A"] not in labels:
                    continue
                missing = [label for label in HEADERS.values() if label not in labels]
                if missing:
                    raise ValueError("Missing required headers: " + "; ".join(missing))
                if any(labels.count(label) != 1 for label in HEADERS.values()):
                    raise ValueError("Duplicate required headers make column mapping ambiguous.")
                mapping = {key: labels.index(label) for key, label in HEADERS.items()}
                candidates.append((sheet.title, rows, index, mapping))
        if len(candidates) != 1:
            raise ValueError(
                "Expected exactly one annualization table with all 46 required headers."
            )
        title, rows, header, mapping = candidates[0]
        employees = []
        totals: dict[str, str] = {}
        for number, row in enumerate(rows[header + 1 :], start=header + 2):
            values = {
                key: text(row[pos].value) if pos < len(row) else "" for key, pos in mapping.items()
            }
            if not any(values.values()):
                continue
            if values["A"].upper() == "TOTAL":
                if totals:
                    raise ValueError("Multiple TOTAL rows are ambiguous.")
                totals = values
                continue
            cells = {key: f"{get_column_letter(pos + 1)}{number}" for key, pos in mapping.items()}
            employees.append(SourceRow(str(number), values, cells))
        if not employees:
            raise ValueError("The workbook has no employee records.")
        heading = " ".join(text(c.value) for row in rows[:header] for c in row if c.value)
        year_match = re.search(r"ANNUALIZATION\s*[—–-]\s*(20\d{2})", heading, re.I)
        name = re.split(r"\s+[—–]\s+", heading)[0] if "ANNUALIZATION" in heading.upper() else ""
        return SourceWorkbook(
            filename,
            hashlib.sha256(data).hexdigest(),
            title,
            employees,
            FilingContext(name=name, year=year_match[1] if year_match else ""),
            totals,
        )
    finally:
        book.close()
