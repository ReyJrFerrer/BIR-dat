import copy
import csv
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from alphalist.conversion import build_snapshot, effective_declarations
from alphalist.dat import CONTROL_INDICES, digest, serialize
from alphalist.domain import REFERENCE, REFERENCE_HASH, WORKBOOK
from alphalist.pdf import render
from alphalist.presentation import data_tasks, grouped_issues
from alphalist.review import audit, correct, load_reference, status
from alphalist.workbook import read_workbook


def test_reference_exact_bytes_and_controls():
    snapshot = build_snapshot(load_reference("validated"))
    data = serialize(snapshot)
    assert data == REFERENCE.read_bytes()
    assert len(data) == 612
    assert digest(data) == REFERENCE_HASH
    rows = list(csv.reader(data.decode().splitlines()))
    assert [len(row) for row in rows] == [4, 49, 36]
    assert data.count(b"\r\n") == 3
    for i, value in zip(CONTROL_INDICES, rows[-1][5:], strict=True):
        assert sum(Decimal(row[i]) for row in rows[1:-1]) == Decimal(value)


def test_original_workbook_has_only_three_data_tasks():
    review = load_reference("workbook")
    snapshot = build_snapshot(review)
    assert len(snapshot.records) == 11
    assert {task.key for task in data_tasks(snapshot)} == {":tin", ":branch", "12:V"}
    assert len(grouped_issues(snapshot, "review")) == 6
    assert len(grouped_issues(snapshot, "profile")) > 0
    assert snapshot.total("I") == Decimal("623085.66")
    assert snapshot.total("O") == Decimal("4100331.50")
    assert snapshot.total("P") == Decimal("422807.50")
    assert snapshot.total("AQ") == Decimal("1050.00")
    assert snapshot.total("AR") == Decimal("423857.50")
    assert snapshot.total("AS") == Decimal("400000.00")
    with pytest.raises(ValueError, match="blocked"):
        serialize(snapshot)
    with pytest.raises(ValueError, match="blocked"):
        render(snapshot)


def test_shared_answers_never_fabricate_values_or_override_exceptions():
    review = load_reference("workbook")
    original = copy.deepcopy(review.source)
    # Classifications only: no missing TIN is invented for any test.
    correct(
        review,
        "filing",
        {"schedule": "non_mwe", "prior": "none"},
        "Test of shared answer precedence",
    )
    assert effective_declarations(review, review.source.rows[0])["schedule"] == "non_mwe"
    carlo = next(r for r in review.source.rows if r.values["A"] == "EMP-0015")
    assert effective_declarations(review, carlo)["prior"] == "present"
    liza = next(r for r in review.source.rows if r.key == "12")
    correct(review, liza.key, {"schedule": "unknown"}, "Test of unresolved employee exception")
    assert effective_declarations(review, liza)["schedule"] == "unknown"
    correct(review, "filing", {"schedule": "mwe"}, "Test of subsequent shared answer")
    assert effective_declarations(review, liza)["schedule"] == "unknown"
    correct(review, liza.key, {"schedule": "inherit"}, "Test of restoring shared answer")
    assert effective_declarations(review, liza)["schedule"] == "mwe"
    assert review.source == original
    assert review.source.rows[6].values["V"] == ""
    assert len(review.history) == 5
    assert not build_snapshot(review).valid
    assert audit(review)["filing_declarations"] == {"schedule": "mwe", "prior": "none"}


def test_user_cannot_dismiss_profile_checks_with_shared_answers():
    review = load_reference("workbook")
    before = {i.code for i in build_snapshot(review).issues if i.category == "profile"}
    correct(
        review,
        "filing",
        {
            "schedule": "non_mwe",
            "prior": "none",
            "withholding": "confirmed",
            "tax_treatment": "confirmed",
            "benefits": "confirmed",
            "substituted": "confirmed",
        },
        "Test-only classification declarations",
    )
    snapshot = build_snapshot(review)
    assert not snapshot.valid
    assert {i.code for i in snapshot.issues if i.category == "profile"} == before
    assert {task.key for task in data_tasks(snapshot)} == {":tin", ":branch", "12:V"}


def test_existing_prior_block_cannot_be_removed_by_declaration():
    review = load_reference("workbook")
    correct(review, "10", {"prior": "none"}, "Test of contradictory declaration")
    assert any(i.code == "PRIOR_CONFLICT" for i in build_snapshot(review).issues)


def test_invalid_answers_do_not_partially_apply():
    review = load_reference("workbook")
    with pytest.raises(ValueError):
        correct(review, "filing", {"schedule": "non_mwe", "prior": "anything"}, "Invalid choices")
    assert not review.history
    assert not review.filing_declarations


def test_corrections_revalidate_and_evidence_follows_hash():
    review = load_reference("validated")
    assert status(review) == "Official validation evidence attached"
    correct(review, "6", {"X": ""}, "Test missing name")
    assert status(review) == "Needs correction"
    assert any(i.field == "X" for i in build_snapshot(review).issues)
    correct(review, "6", {"X": "Maria"}, "Restore the original supplied name")
    assert status(review) == "Official validation evidence attached"
    correct(review, "6", {"AS": "28999.00"}, "Test adjusted amount recalculation")
    assert status(review) == "Passed internal checks"
    assert not audit(review)["evidence_matches_current_output"]


def test_prior_and_refund_amounts_are_preserved_in_draft():
    snapshot = build_snapshot(load_reference("workbook"))
    rows = {r.employee_id: r for r in snapshot.records}
    assert rows["EMP-0015"].amount("combined_taxable") == Decimal("573800")
    assert rows["EMP-0015"].amount("AR") == Decimal("57260")
    assert rows["EMP-0011"].amount("refund") == Decimal("3825")
    assert rows["EMP-0011"].amount("additional") == 0
    assert "Ñ" in rows["EMP-0004"].name


def test_pdf_reconciles_and_draft_accounts_for_every_employee():
    snapshot = build_snapshot(load_reference("validated"))
    pdf = PdfReader(io.BytesIO(render(snapshot)))
    assert len(pdf.pages) == 1
    assert tuple(pdf.pages[0].mediabox) == (0, 0, 1008, 612)
    text = pdf.pages[0].extract_text()
    assert "449,812.50" in text and "32,462.50" in text
    assert "Grand total" in text and "END OF REPORT" in text
    snapshot = build_snapshot(load_reference("workbook"))
    draft = PdfReader(io.BytesIO(render(snapshot, draft=True)))
    assert len(draft.pages) > 1
    texts = [page.extract_text() for page in draft.pages]
    for text in texts:
        assert "DRAFT" in text and "Page subtotal" in text
    for row in snapshot.records:
        assert row.employee_id in "\n".join(texts)
    assert "Grand total" in texts[-1]


def test_moved_headers_and_columns_keep_source_coordinates():
    book = load_workbook(WORKBOOK)
    sheet = book.active
    sheet.insert_rows(1, 2)
    sheet.insert_cols(1)
    buffer = io.BytesIO()
    book.save(buffer)
    result = read_workbook(buffer.getvalue(), "moved.xlsx")
    assert len(result.rows) == 11
    assert result.rows[6].cells["V"] == "W14"
    assert result.rows[6].values["V"] == ""


@pytest.mark.parametrize("kind", ["duplicate", "missing"])
def test_ambiguous_headers_rejected(kind):
    book = load_workbook(WORKBOOK)
    sheet = book.active
    if kind == "duplicate":
        sheet["AU5"] = sheet["V5"].value
    else:
        sheet["V5"] = "wrong header"
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(ValueError, match="headers"):
        read_workbook(buffer.getvalue(), "bad.xlsx")
